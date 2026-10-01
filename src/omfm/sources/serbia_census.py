"""Serbia, Census 2022: daily migrants by municipality (SORS, Excel tables).

What the files do not contain: an origin x destination matrix. Destinations are four nested
administrative rings,

    total
      same area (oblast)
        same municipality          <- numerator of self-containment
        another municipality
      another area
      abroad

so self-containment is computable and delimitation is not.

Definition that drives the results: a *daily migrant* leaves the SETTLEMENT of residence, not
the municipality. A municipality that consists of one settlement therefore has
self-containment exactly 0: every one of its migrants has left it by definition.

Sheet layout: a hierarchy in the first column. Upper-case rows are the republic and
macro-regions; "регион" rows are regions; "област" rows are areas; "Град X" rows inside an
area are aggregates over city municipalities (not municipalities); the rest are municipalities.
Every territory is followed by settlement-type rows ("Градска" / "Остала") and by sex
("с" total, "м", "ж").
"""

from __future__ import annotations

import pandas as pd

from ..analysis.selfcontainment import band_self_containment

__all__ = ["read_daily_migration", "add_self_containment", "COLS"]

SEX_TOTAL = "с"
_SETTLEMENT_ROWS = {"Градска", "Остала"}
COLS = {2: "total", 3: "same_municipality", 4: "other_municipality_same_area", 5: "other_area", 6: "abroad"}


def _classify(name: str) -> str:
    low = name.lower()
    if name.isupper():
        return "country_or_macroregion"
    if "регион" in low:
        return "region"
    if "област" in low:
        return "area"
    if name.startswith("Град "):
        return "city_aggregate"
    return "municipality"


def read_daily_migration(path, sheet: str | None = None, level: str = "municipality") -> pd.DataFrame:
    """Parse one SORS workbook.

    ``level="municipality"`` returns municipalities only; ``level="all"`` keeps every row of the
    hierarchy with a ``level`` column (used by the reconciliation test against the republic total).
    ``single_settlement`` is derived from the file: the municipality has a "Градска" row and no
    "Остала" row.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    recs, cur_area, cur_city, cur = [], None, None, None
    for row in ws.iter_rows(min_row=5, values_only=True):
        raw, sex = row[0], str(row[1]).strip()
        if raw is None or not str(raw).strip():
            continue
        name = str(raw).strip()
        if name in _SETTLEMENT_ROWS:
            if cur is not None and sex == SEX_TOTAL:
                cur["settlement_types"].add(name)
            continue
        kind = _classify(name)
        if kind == "area":
            cur_area, cur_city = name, None
        elif kind == "city_aggregate":
            cur_city = name
        if sex != SEX_TOTAL:
            continue
        rec = {
            "name": name,
            "level": kind,
            "area": cur_area,
            "city": cur_city if kind == "municipality" else None,
            "settlement_types": set(),
        }
        try:
            for idx, col in COLS.items():
                rec[col] = float(row[idx]) if row[idx] is not None else 0.0
        except (TypeError, ValueError):
            cur = None
            continue
        recs.append(rec)
        cur = rec
    df = pd.DataFrame(recs)
    df["single_settlement"] = df.settlement_types.apply(lambda t: t == {"Градска"})
    df = df.drop(columns="settlement_types")
    if level == "municipality":
        df = df[df.level == "municipality"].drop(columns="level")
    return df.reset_index(drop=True)


def add_self_containment(df: pd.DataFrame) -> pd.DataFrame:
    """Add ``sc`` = migrants staying in their municipality / all daily migrants, and a join key."""
    out = df.copy()
    out["sc"] = band_self_containment(out.total, out.same_municipality)
    out["key"] = out.area.fillna("") + " / " + out.name
    return out
